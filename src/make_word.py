"""Build word/paper.docx (same text, figures, tables, captions, references as latex/main.tex)
and word/figures_for_submission.pptx (one figure per slide).

Run after the LaTeX build (needs latex/main.bbl for reference numbering):
    python code/src/make_word.py
"""
import re

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt

from common import PAPER, ROOT

WORD = ROOT.parent / "word"


# ---------------------------------------------------------------- helpers
def load_macros():
    m = {}
    for line in (PAPER / "numbers.tex").read_text(encoding="utf8").splitlines():
        g = re.match(r"\\newcommand\{\\(\w+)\}\{(.*)\}$", line)
        if g:
            m[g.group(1)] = g.group(2)
    return m


def bbl_entries():
    """Return ordered list of (key, plain-text reference) from main.bbl."""
    t = (PAPER / "main.bbl").read_text(encoding="utf8")
    items = re.split(r"\\bibitem\{", t)[1:]
    out = []
    for it in items:
        key, body = it.split("}", 1)
        body = body.split("\\end{thebibliography}")[0]
        body = re.sub(r"\\BIBentryALTinterwordspacing|\\BIBentrySTDinterwordspacing", "", body)
        body = re.sub(r"\\newblock", " ", body)
        body = re.sub(r"\\hskip\s*[\d.]+em\s*plus\s*[\d.]+em\s*minus\s*[\d.]+em\\relax", " ", body)
        body = re.sub(r"[\d.]+em plus [\d.]+em minus [\d.]+em", " ", body)
        body = re.sub(r"\\url\{([^}]*)\}", r"\1", body)
        body = re.sub(r"\\href\{[^}]*\}\{([^}]*)\}", r"\1", body)
        out.append((key, " ".join(latex_plain(body).split())))
    return out


ACCENTS = {r"\'c": "ć", r"\'o": "ó", r'\"o': "ö", r"\'e": "é", r'\"u': "ü"}


def latex_plain(s):
    for k, v in ACCENTS.items():
        s = s.replace("{" + k + "}", v).replace(k, v)
    s = s.replace("``", "\u201c").replace("''", "\u201d").replace("---", "\u2014").replace("--", "\u2013")
    s = s.replace("\\&", "&").replace("\\%", "%").replace("~", " ").replace("\\ ", " ")
    s = re.sub(r"\\emph\{([^}]*)\}|\\textit\{([^}]*)\}|\\textbf\{([^}]*)\}",
               lambda g: g.group(1) or g.group(2) or g.group(3), s)
    s = re.sub(r"\\[a-zA-Z]+\*?", "", s)
    s = s.replace("{", "").replace("}", "")
    return s


MATH = [(r"\mathrm{PUE}", "PUE"), (r"\widehat{\mathrm{CI}}", "CI\u0302"), (r"\epsilon", "\u03b5"),
        (r"\times", "\u00d7"), (r"\in", "\u2208"), (r"\le", "\u2264"), (r"\mathrm{CI}", "CI")]


def math_runs(s):
    """Convert a short inline-math string to (text, sub, sup, italic) runs."""
    for a, b in MATH:
        s = s.replace(a, b)
    runs = []
    for tok in re.split(r"(_\{[^}]*\}|_.|\^\{[^}]*\})", s):
        if not tok:
            continue
        if tok.startswith("_"):
            runs.append((tok[1:].strip("{}"), "sub"))
        elif tok.startswith("^"):
            runs.append((tok[1:].strip("{}").replace("*", "*"), "sup"))
        else:
            txt = tok.replace("{", "").replace("}", "").replace("\\", "")
            runs.append((txt, "ital" if re.fullmatch(r"[A-Za-z]", txt.strip()) else None))
    return runs


class Ctx:
    def __init__(self, macros, cites, figs, tabs):
        self.macros, self.cites, self.figs, self.tabs = macros, cites, figs, tabs


def add_text(par, s, ctx, bold=False, italic=False, size=None):
    """Add LaTeX-ish inline text to a python-docx paragraph."""
    s = re.sub(r"\\(\w+)\{\}", lambda g: ctx.macros.get(g.group(1), g.group(0)), s)
    s = re.sub(r"\\(" + "|".join(sorted(ctx.macros, key=len, reverse=True)) + r")(?![A-Za-z])",
               lambda g: ctx.macros[g.group(1)], s)
    s = re.sub(r"~?\\cite\{([^}]*)\}",
               lambda g: " " + "".join(f"[{ctx.cites[k.strip()]}]" for k in g.group(1).split(",")), s)
    s = re.sub(r"\\ref\{([^}]*)\}", lambda g: str(ctx.figs.get(g.group(1), ctx.tabs.get(g.group(1), "?"))), s)
    s = re.sub(r"\\eqref\{[^}]*\}", "(1)", s)
    s = s.replace("\\ldots", "\u2026").replace("a.m.\\", "a.m.").replace("\\,", " ")
    # tokenise: \textbf{..}, \emph{..}, \textit{..}, $..$
    pos = 0
    pat = re.compile(r"\\textbf\{([^{}]*)\}|\\emph\{([^{}]*)\}|\\textit\{([^{}]*)\}|\$([^$]*)\$")
    for g in pat.finditer(s):
        _plain(par, s[pos:g.start()], bold, italic, size)
        if g.group(1) is not None:
            add_text(par, g.group(1), ctx, True, italic, size)
        elif g.group(2) is not None or g.group(3) is not None:
            add_text(par, g.group(2) or g.group(3), ctx, bold, True, size)
        else:
            for txt, kind in math_runs(g.group(4)):
                r = par.add_run(txt)
                r.bold, r.italic = bold, (kind == "ital") or italic
                r.font.subscript = kind == "sub"
                r.font.superscript = kind == "sup"
                if size:
                    r.font.size = Pt(size)
        pos = g.end()
    _plain(par, s[pos:], bold, italic, size)


def _plain(par, s, bold, italic, size):
    if not s:
        return
    r = par.add_run(latex_plain(s))
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)


def parse_tabular(path):
    rows = []
    t = path.read_text(encoding="utf8")
    body = t.split("\n", 1)[1].rsplit("\\end{tabular}", 1)[0]
    for line in body.split("\\\\"):
        line = line.replace("\\hline", "").strip()
        if not line:
            continue
        cells = []
        for c in line.split("&"):
            c = c.strip()
            m = re.match(r"\\multicolumn\{(\d+)\}\{[^}]*\}\{(.*)\}$", c)
            if m:
                cells += [m.group(2)] + [""] * (int(m.group(1)) - 1)
            else:
                cells.append(c)
        rows.append(cells)
    return rows


def equation_png():
    out = WORD / "eq1.png"
    fig = plt.figure(figsize=(3.6, 0.45))
    fig.text(0.02, 0.35, r"$C(s,r)=\frac{E}{d}\,(1+\epsilon_r)\,\mathrm{PUE}_r\sum_{h=s}^{s+d-1}"
             r"\widehat{\mathrm{CI}}_r(h)$", fontsize=12)
    fig.savefig(out, dpi=300, bbox_inches="tight", transparent=True)
    plt.close(fig)
    return out


# ---------------------------------------------------------------- main
def build():
    WORD.mkdir(exist_ok=True)
    tex = (PAPER / "main.tex").read_text(encoding="utf8")
    tex = re.sub(r"(?<!\\)%.*", "", tex)
    body = tex.split("\\begin{document}", 1)[1].split("\\end{document}", 1)[0]
    macros = load_macros()
    refs = bbl_entries()
    cites = {k: i + 1 for i, (k, _) in enumerate(refs)}
    figs = {lab: i + 1 for i, lab in enumerate(re.findall(r"\\begin\{figure\*?\}.*?\\label\{([^}]*)\}", body, re.S))}
    tabs = {lab: i + 1 for i, lab in enumerate(re.findall(r"\\begin\{table\}.*?\\label\{([^}]*)\}", body, re.S))}
    ctx = Ctx(macros, cites, figs, tabs)

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(10.5)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(0.9)

    title = re.search(r"\\title\{(.*?)\}\n", body, re.S).group(1)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, title, ctx, bold=True, size=18)
    # author block, read from main.tex: "Name, \textit{Institution}\\ ..." followed by \thanks{...}
    ablock = body.split("\\author{", 1)[1].split("\\markboth", 1)[0]
    head, *thanks = ablock.split("\\thanks{")
    for line in head.split("\\\\"):
        line = " ".join(line.replace("%", "").split())
        if not line:
            continue
        name, _, inst = line.partition(",")
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        add_text(p, name.strip(), ctx, size=12)
        add_text(p, ", " + inst.strip(), ctx, size=10)
    for t in thanks:
        t = " ".join(t.replace("%", "").split()).rstrip("}").strip()
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        add_text(p, t, ctx, size=8.5)
    abstract = re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}", body, re.S).group(1)
    p = doc.add_paragraph()
    add_text(p, "Abstract\u2014", ctx, bold=True, italic=True)
    add_text(p, " ".join(abstract.split()), ctx, bold=True)

    main = body.split("\\end{abstract}", 1)[1].split("\\bibliographystyle", 1)[0]
    blocks = re.split(r"(\\section\{[^}]*\}|\\begin\{figure\*?\}.*?\\end\{figure\*?\}|"
                      r"\\begin\{table\}.*?\\end\{table\}|\\begin\{equation\}.*?\\end\{equation\})", main, flags=re.S)
    fig_files = []
    for b in blocks:
        if not b or not b.strip():
            continue
        if b.startswith("\\section"):
            h = doc.add_heading(level=1)
            add_text(h, re.match(r"\\section\{(.*)\}", b).group(1), ctx)
            for r in h.runs:
                r.font.color.rgb = RGBColor(0, 0, 0); r.font.name = "Times New Roman"; r.font.size = Pt(12)
        elif b.startswith("\\begin{figure"):
            f = re.search(r"\\includegraphics\[[^]]*\]\{([^}]*)\}", b).group(1)
            png = PAPER / (f + ".png")
            fig_files.append(png)
            doc.add_picture(str(png), width=Inches(6.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            lab = re.search(r"\\label\{([^}]*)\}", b).group(1)
            cap = re.search(r"\\caption\{(.*)\}\s*\\label", b, re.S).group(1)
            p = doc.add_paragraph()
            add_text(p, f"FIGURE {figs[lab]}. ", ctx, bold=True, size=9)
            add_text(p, " ".join(cap.split()), ctx, size=9)
        elif b.startswith("\\begin{table"):
            lab = re.search(r"\\label\{([^}]*)\}", b).group(1)
            cap = re.search(r"\\caption\{(.*)\}\s*\\label", b, re.S).group(1)
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_text(p, f"TABLE {tabs[lab]}. ", ctx, bold=True, size=9)
            add_text(p, " ".join(cap.split()), ctx, size=9)
            src = re.search(r"\\input\{([^}]*)\}", b).group(1)
            rows = parse_tabular(PAPER / (src + ".tex"))
            tbl = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            tbl.style = "Table Grid"
            for i, row in enumerate(rows):
                for j, c in enumerate(row):
                    cell = tbl.cell(i, j)
                    cell.text = ""
                    add_text(cell.paragraphs[0], c, ctx, bold=(i == 0 or (i == 1 and "Region" in row[0] or "Strategy" in row[0])), size=8.5)
            doc.add_paragraph()
        elif b.startswith("\\begin{equation"):
            doc.add_picture(str(equation_png()), width=Inches(3.4))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.paragraphs[-1].add_run("        (1)")
        else:
            for para in re.split(r"\n\s*\n", b):
                para = " ".join(para.split())
                if not para:
                    continue
                para = re.sub(r"\\chapterinitial\{([^}]*)\}", r"\\textbf{\1}", para)
                if para.startswith("\\begin{IEEEbio") or para.startswith("\\maketitle"):
                    continue
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(4)
                add_text(p, para, ctx)
    h = doc.add_heading("REFERENCES", level=1)
    for r in h.runs:
        r.font.color.rgb = RGBColor(0, 0, 0); r.font.name = "Times New Roman"; r.font.size = Pt(12)
    for i, (_, txt) in enumerate(refs):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        add_text(p, f"[{i + 1}] {txt}", ctx, size=9)
    for bio in re.finditer(r"\\begin\{IEEEbiographynophoto\}\{([^}]*)\}(.*?)\\end\{IEEEbiographynophoto\}",
                           body, re.S):
        p = doc.add_paragraph()
        add_text(p, bio.group(1) + " ", ctx, bold=True)
        add_text(p, " ".join(bio.group(2).split()), ctx)
    doc.save(WORD / "paper.docx")
    (WORD / "eq1.png").unlink(missing_ok=True)

    # PowerPoint: one figure per slide (16:9), caption in the notes
    prs = Presentation()
    prs.slide_width, prs.slide_height = PInches(13.333), PInches(7.5)
    caps = re.findall(r"\\caption\{(.*?)\}\s*\\label\{(fig:[^}]*)\}", body, re.S)
    for png, (cap, lab) in zip(fig_files, caps):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tb = s.shapes.add_textbox(PInches(0.4), PInches(0.2), PInches(12.5), PInches(0.6))
        tb.text_frame.text = f"Figure {figs[lab]}"
        tb.text_frame.paragraphs[0].runs[0].font.size = PPt(24)
        tb.text_frame.paragraphs[0].runs[0].font.bold = True
        pic = s.shapes.add_picture(str(png), PInches(0.4), PInches(1.0), width=PInches(12.5))
        if pic.height > PInches(6.3):
            ratio = PInches(6.3) / pic.height
            pic.height = PInches(6.3); pic.width = int(pic.width * ratio)
        s.notes_slide.notes_text_frame.text = f"FIGURE {figs[lab]}. " + latex_plain(
            re.sub(r"~?\\cite\{([^}]*)\}", lambda g: " " + "".join(f"[{cites[k.strip()]}]" for k in g.group(1).split(",")),
                   " ".join(cap.split())).replace("$", ""))
    prs.save(WORD / "figures_for_submission.pptx")
    print("wrote", WORD / "paper.docx", "and", WORD / "figures_for_submission.pptx")


if __name__ == "__main__":
    build()
